# -*- coding: utf-8 -*-
"""Print the dataset itself -- the legal text and nothing else.

    python -m report.build_dataset_plain

Writes report/ชุดข้อมูลจรรยาบรรณวิชาชีพครู.docx

Companion to report.build_dataset_report_ksp, which documents how the dataset was
built and is full of script names, field names and JSON. This one carries no
technical material at all: the ten documents, what each one covers, the five duties
set out in full, and then every ข้อ and มาตรา in the corpus verbatim. A reader who
only wants to check the content against the Royal Gazette reads this file.

Every word of legal text comes straight out of data/processed/corpus_ksp.jsonl, so
rebuilding after a corpus change keeps the two in step.
"""
from __future__ import annotations

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

from report.build_dataset_report_ksp import (      # noqa: E402
    DUTIES, ORDER, PROFESSIONS, load, by_doc, sec_key, sources, th,
)

OUT = os.path.join(HERE, "ชุดข้อมูลจรรยาบรรณวิชาชีพครู.docx")

# Smaller than the companion report. That one is prose and tables; this one is a
# wall of legal text, where 14 pt reads better and fits far more to a page.
BODY_PT = 14
HEAD_PT = 14
TITLE_PT = 18

# Paragraph spacing, in points. Legal text is a run of short numbered items, so the
# gap between them does the work that indentation does in prose -- too little and
# the items run together, too much and one ข้อ eats a third of a page.
SPACE_AFTER = 4
SPACE_BEFORE_HEAD = 12


def main() -> None:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    records = load()
    docs = by_doc(records)
    src = sources()

    doc = Document()

    def style(name, size, bold, space_before=0, space_after=SPACE_AFTER):
        st = doc.styles[name]
        st.font.name = "Sarabun"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = RGBColor(0, 0, 0)
        st.element.rPr.rFonts.set(qn("w:cs"), "Sarabun")
        st.element.rPr.rFonts.set(qn("w:eastAsia"), "Sarabun")
        st.element.rPr.rFonts.set(qn("w:ascii"), "Sarabun")
        st.element.rPr.rFonts.set(qn("w:hAnsi"), "Sarabun")
        # Word keeps a separate size for complex scripts; Thai runs follow that one,
        # not w:sz, so setting only the Latin size leaves Thai at the default.
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        st.element.rPr.append(szcs)
        pf = st.paragraph_format
        pf.space_before = Pt(space_before)
        pf.space_after = Pt(space_after)
        pf.line_spacing = 1.0
        pf.widow_control = True
        return st

    style("Normal", BODY_PT, None)
    style("Heading 1", HEAD_PT, True, space_before=SPACE_BEFORE_HEAD)
    style("Heading 2", HEAD_PT, True, space_before=SPACE_BEFORE_HEAD)
    style("Heading 3", HEAD_PT, True, space_before=8)
    style("List Bullet", BODY_PT, None)

    def fix(run, size=BODY_PT, bold=None, italic=None):
        run.font.name = "Sarabun"
        for attr in ("w:cs", "w:ascii", "w:hAnsi", "w:eastAsia"):
            run._element.rPr.rFonts.set(qn(attr), "Sarabun")
        run.font.size = Pt(size)
        szcs = OxmlElement("w:szCs")
        szcs.set(qn("w:val"), str(int(size * 2)))
        run._element.rPr.append(szcs)
        if bold is not None:
            run.font.bold = bold
            b = OxmlElement("w:bCs")
            b.set(qn("w:val"), "1" if bold else "0")
            run._element.rPr.append(b)
        if italic:
            run.font.italic = True
        return run

    def para(text="", style_name=None, size=BODY_PT, bold=None, center=False,
             space_after=None):
        p = doc.add_paragraph(style=style_name)
        if text:
            fix(p.add_run(text), size=size, bold=bold)
        if center:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if space_after is not None:
            p.paragraph_format.space_after = Pt(space_after)
        return p

    def lead(label, body, size=BODY_PT):
        """One paragraph that opens in bold and continues in normal weight."""
        p = doc.add_paragraph()
        fix(p.add_run(label), size=size, bold=True)
        if body:
            fix(p.add_run("  " + body), size=size)
        return p

    def table(rows, widths=None, head=True):
        t = doc.add_table(rows=0, cols=len(rows[0]))
        t.style = "Table Grid"
        for n, row in enumerate(rows):
            cells = t.add_row().cells
            for j, value in enumerate(row):
                cells[j].text = ""
                p = cells[j].paragraphs[0]
                p.paragraph_format.space_after = Pt(1)
                p.paragraph_format.space_before = Pt(1)
                fix(p.add_run(str(value)), size=BODY_PT - 1,
                    bold=True if (head and n == 0) else None)
        return t

    # ------------------------------------------------------------------ cover
    para("ชุดข้อมูลจรรยาบรรณของวิชาชีพครู", size=TITLE_PT, bold=True, center=True,
         space_after=2)
    para("ตัวบทที่ใช้ในระบบถาม-ตอบ ฉบับเต็ม", size=BODY_PT, center=True, space_after=14)

    n_rec = len(records)
    n_units = len({(r["sysid"], r["section"]) for r in records})
    n_chars = sum(r["n_chars"] for r in records)
    n_kho = sum(1 for r in records if r["unit"] == "ข้อ")
    n_matra = n_rec - n_kho

    table([
        ["รายการ", "จำนวน"],
        ["เอกสาร", f"{len(docs)} ฉบับ"],
        ["ข้อและมาตรา", f"{n_units}"],
        ["ระเบียนข้อความ", f"{n_rec}"],
        ["นับเป็น “ข้อ”", f"{n_kho}"],
        ["นับเป็น “มาตรา”", f"{n_matra}"],
        ["ความยาวรวม", f"{n_chars:,} ตัวอักษร"],
        ["ด้านจรรยาบรรณ", f"{len(DUTIES)} ด้าน"],
        ["วิชาชีพที่ครอบคลุม", f"{len(PROFESSIONS)} ประเภท"],
    ])
    doc.add_page_break()

    # ------------------------------------------------------------------ toc
    para("สารบัญ", style_name="Heading 1")
    rows = [["ตอน", "เรื่อง"],
            ["๑", "รายการเอกสารทั้งหมด"],
            ["๒", "จรรยาบรรณของวิชาชีพ ๕ ด้าน"],
            ["๓", "แบบแผนพฤติกรรมตามจรรยาบรรณ แยกตามวิชาชีพ ๔ ประเภท"]]
    for i, key in enumerate(ORDER, 4):
        rows.append([th(i), docs[key][0]["short"]])
    table(rows)
    para()
    para("สารบัญอัตโนมัติ (คลิกขวาแล้วเลือก Update Field)", bold=True,
         space_after=2)
    p = doc.add_paragraph()
    r = p.add_run()
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                             ("w:instrText", {"xml:space": "preserve"},
                              r' TOC \o "1-2" \h \z \u '),
                             ("w:fldChar", {"w:fldCharType": "separate"}, None),
                             ("w:t", {}, "—"),
                             ("w:fldChar", {"w:fldCharType": "end"}, None)):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text is not None:
            el.text = text
        r._r.append(el)
    doc.add_page_break()

    # "ตอน" not "ส่วน": the documents use "ส่วนที่ N" for their own subdivisions
    # ------------------------------------------------------------------ part 1
    para("ตอนที่ ๑ รายการเอกสารทั้งหมด", style_name="Heading 1")
    rows = [["#", "ชื่อเอกสาร", "หน่วย", "จำนวน", "วันประกาศ", "สถานะ"]]
    for i, key in enumerate(ORDER, 1):
        rs = docs[key]
        units = len({r["section"] for r in rs})
        rows.append([i, rs[0]["short"], rs[0]["unit"], units,
                     rs[0]["published"] or "—",
                     "ยกเลิกแล้ว" if rs[0]["superseded_by"] else "ใช้บังคับ"])
    table(rows)

    para()
    para("ที่มาของแต่ละฉบับ", style_name="Heading 2")
    for key in ORDER:
        rs = docs[key]
        lead(rs[0]["short"], "")
        para(rs[0]["act_full"], space_after=1)
        para(src[key][1], size=BODY_PT - 2, space_after=6)
    doc.add_page_break()

    # ------------------------------------------------------------------ part 2
    para("ตอนที่ ๒ จรรยาบรรณของวิชาชีพ ๕ ด้าน", style_name="Heading 1")
    para("ตามข้อบังคับคุรุสภา ว่าด้วยจรรยาบรรณของวิชาชีพ พ.ศ. ๒๕๕๖",
         space_after=8)

    rows = [["ด้าน", "ข้อบังคับ ๒๕๕๖"] + [f"๒๕๕๐ ({p})" for p in PROFESSIONS]]
    for duty in DUTIES:
        cells = [duty]
        hits = sorted({r["section"] for r in records
                       if r["sysid"] == "ksp-2556" and r["ethics_category"] == duty},
                      key=lambda s: int(s))
        cells.append("ข้อ " + ", ".join(hits))
        for prof in PROFESSIONS:
            got = sorted({r["section"] for r in records
                          if r["sysid"] == "ksp-2550" and r["ethics_category"] == duty
                          and any(prof in c for c in r["chapters"])},
                         key=lambda s: int(s))
            cells.append("ข้อ " + ", ".join(got) if got else "—")
        rows.append(cells)
    table(rows)
    para()

    for duty in DUTIES:
        para(f"จรรยาบรรณ{duty}", style_name="Heading 2")
        hits = sorted([r for r in records if r["sysid"] == "ksp-2556"
                       and r["ethics_category"] == duty], key=sec_key)
        for r in hits:
            head = f"{r['unit']} {r['section']}"
            if r["part"]:
                head += f" (ต่อ)"
            lead(head, r["text"])
    doc.add_page_break()

    # ------------------------------------------------------------------ part 3
    para("ตอนที่ ๓ แบบแผนพฤติกรรมตามจรรยาบรรณ แยกตามวิชาชีพ", style_name="Heading 1")
    para("ตามข้อบังคับคุรุสภา ว่าด้วยแบบแผนพฤติกรรมตามจรรยาบรรณของวิชาชีพ พ.ศ. ๒๕๕๐",
         space_after=8)
    for prof in PROFESSIONS:
        para(f"วิชาชีพ{prof}", style_name="Heading 2")
        hits = sorted([r for r in records if r["sysid"] == "ksp-2550"
                       and any(prof in c for c in r["chapters"])], key=sec_key)
        if not hits:
            raise SystemExit(f"no records for profession {prof}")
        seen = None
        for r in hits:
            if r["ethics_category"] != seen:
                para(f"จรรยาบรรณ{r['ethics_category']}", style_name="Heading 3")
                seen = r["ethics_category"]
            head = f"{r['unit']} {r['section']}"
            if r["part"]:
                head += " (ต่อ)"
            lead(head, r["text"])
    doc.add_page_break()

    # ------------------------------------------------------------------ part 4+
    for i, key in enumerate(ORDER, 4):
        rs = sorted(docs[key], key=sec_key)
        para(f"ตอนที่ {th(i)} {rs[0]['short']}", style_name="Heading 1")
        para(rs[0]["act_full"], space_after=2)
        note = f"{len({r['section'] for r in rs})} {rs[0]['unit']}"
        if rs[0]["published"]:
            note += f" · ประกาศ {rs[0]['published']}"
        if rs[0]["superseded_by"]:
            note += " · ยกเลิกแล้ว"
        para(note, size=BODY_PT - 1, space_after=8)

        last = None
        for r in rs:
            if r["chapters"] != last:
                for c in r["chapters"]:
                    para(c, style_name="Heading 2")
                last = r["chapters"]
            head = f"{r['unit']} {r['section']}"
            if r["part"]:
                head += " (ต่อ)"
            lead(head, r["text"])
        if i - 3 < len(ORDER):
            doc.add_page_break()

    doc.save(OUT)
    print(OUT)
    print(f"{len(docs)} ฉบับ / {n_units} ข้อและมาตรา / {n_rec} ระเบียน / {n_chars:,} ตัวอักษร")


if __name__ == "__main__":
    main()
