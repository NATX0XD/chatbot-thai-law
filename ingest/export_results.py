# -*- coding: utf-8 -*-
"""Write the whole test run out as one Excel workbook.

    .venv/bin/python -m ingest.export_results 6

Five sheets:

  ผลการทดสอบ     one row per question: what was asked, what the system answered,
                 which rules it cited, which rules were expected, and how each
                 of the three assessors coded it, with their reason
  สรุปรายบุคคล    tables 4-1 and 4-2 of the thesis, recomputed from the codings
  สรุปรายรอบ      every round measured so far, per assessor
  ชุดข้อมูล       the instruments the corpus holds, with source and date
  เกณฑ์การให้รหัส  what each outcome code means and what counts as incorrect

The expectation column is not hand-written. Where all three assessors coded an
answer correct, the rules that answer cited are the expected rules -- three
independent readers checked them against the text. Where they did not, the
expectation is whichever rule their notes name, and blank when the notes name
none. Guessing would put a number in the thesis that nobody had verified.
"""
from __future__ import annotations

import json
import os
import re
import sys
from collections import Counter

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVAL = os.path.join(BASE_DIR, "data", "eval")
CORPUS = os.path.join(BASE_DIR, "data", "processed", "corpus_ksp.jsonl")

CORRECT = {1, 2}
CODE_LABEL = {
    1: "1 ถามตรงตัวบท ตอบถูก",
    2: "2 ถามภาษาพูด ตอบถูก",
    0: "0 ถามตรงตัวบท ตอบไม่ถูก",
    3: "3 ถามภาษาพูด ตอบไม่ถูก",
}
CITATION = re.compile(r"(ข้อ|มาตรา)\s*([๐-๙0-9]{1,3}(?:/[๐-๙0-9]{1,3})?)"
                      r"((?:\s*\([ก-ฮ๐-๙0-9]{1,3}\))*)")
THAI_DIGITS = str.maketrans("๐๑๒๓๔๕๖๗๘๙", "0123456789")

HEAD = Font(bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="4F6228")
WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")
RIGHT_FILL = PatternFill("solid", fgColor="E8F0DE")
WRONG_FILL = PatternFill("solid", fgColor="F8DCDC")
SPLIT_FILL = PatternFill("solid", fgColor="FFF4D6")


def load(path: str):
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def rounds_available() -> list[int]:
    return sorted(int(m.group(1)) for m in
                  (re.match(r"coding_round(\d+)_assessor1\.json$", f)
                   for f in os.listdir(EVAL)) if m)


def coding(round_no: int) -> dict[int, dict[str, dict]]:
    out = {}
    for n in (1, 2, 3):
        path = os.path.join(EVAL, f"coding_round{round_no}_assessor{n}.json")
        if os.path.exists(path):
            out[n] = {r["id"]: r for r in load(path)}
    return out


def cited(answer: str) -> list[str]:
    """The citations an answer prints, deduplicated, in order."""
    seen, out = set(), []
    for m in CITATION.finditer(answer):
        one = f"{m.group(1)} {m.group(2).translate(THAI_DIGITS)}" \
              f"{''.join(m.group(3).split())}"
        if one not in seen:
            seen.add(one)
            out.append(one)
    return out


def expected(row_id: str, answer: str, codes: dict[int, dict]) -> tuple[str, str]:
    """(the rules expected, where that expectation came from)."""
    marks = [c[row_id]["code"] for c in codes.values() if row_id in c]
    if marks and all(m in CORRECT for m in marks):
        return ", ".join(cited(answer)), "ผู้ประเมินทั้งสามคนตรวจแล้วว่าถูก"
    named = []
    for n, c in codes.items():
        if row_id in c and c[row_id]["code"] not in CORRECT:
            for m in CITATION.finditer(c[row_id].get("note", "")):
                one = f"{m.group(1)} {m.group(2).translate(THAI_DIGITS)}" \
                      f"{''.join(m.group(3).split())}"
                if one not in named:
                    named.append(one)
    if named:
        return ", ".join(named), "จากบันทึกเหตุผลของผู้ประเมินที่ให้ไม่ถูก"
    return "", "ยังไม่ได้ยืนยันกับตัวบท"


def widths(sheet, sizes: dict[str, int]) -> None:
    for col, size in sizes.items():
        sheet.column_dimensions[col].width = size


def header(sheet, names: list[str]) -> None:
    sheet.append(names)
    for cell in sheet[1]:
        cell.font, cell.fill, cell.alignment = HEAD, HEAD_FILL, WRAP
    sheet.freeze_panes = "A2"


def sheet_results(book, round_no: int) -> None:
    answers = {r["id"]: r for r in
               load(os.path.join(EVAL, f"user_answers_round{round_no}.json"))}
    codes = coding(round_no)
    sheet = book.create_sheet(f"ผลการทดสอบ รอบ {round_no}")
    header(sheet, [
        "รหัสคำถาม", "หมวดจรรยาบรรณ", "ลักษณะคำถาม", "คำถาม",
        "คำตอบของระบบ", "ตัวบทที่ระบบอ้าง", "ตัวบทที่คาดหวัง",
        "ที่มาของความคาดหวัง",
        "คนที่ 1", "เหตุผลคนที่ 1", "คนที่ 2", "เหตุผลคนที่ 2",
        "คนที่ 3", "เหตุผลคนที่ 3", "สรุป", "เวลา (วินาที)",
    ])
    for row_id, a in answers.items():
        marks = [codes[n][row_id]["code"] for n in sorted(codes)
                 if row_id in codes[n]]
        right = sum(1 for m in marks if m in CORRECT)
        verdict = ("ยังไม่ได้ตรวจ" if not marks
                   else "ถูก" if right == len(marks)
                   else "ไม่ถูก" if right == 0
                   else f"เห็นไม่ตรงกัน ({right}/{len(marks)} ว่าถูก)")
        want, source = expected(row_id, a["answer"], codes)
        line = [row_id, a["category"], a["phrasing"], a["question"],
                a["answer"], "\n".join(cited(a["answer"])), want, source]
        for n in (1, 2, 3):
            entry = codes.get(n, {}).get(row_id)
            line += [CODE_LABEL.get(entry["code"], "") if entry else "",
                     entry.get("note", "") if entry else ""]
        line += [verdict, a.get("seconds", "")]
        sheet.append(line)
        fill = (RIGHT_FILL if verdict == "ถูก" else
                WRONG_FILL if verdict == "ไม่ถูก" else
                SPLIT_FILL if marks else None)
        for cell in sheet[sheet.max_row]:
            cell.alignment = WRAP
            if fill:
                cell.fill = fill
    widths(sheet, {"A": 10, "B": 22, "C": 12, "D": 42, "E": 70, "F": 20,
                   "G": 20, "H": 26, "I": 18, "J": 40, "K": 18, "L": 40,
                   "M": 18, "N": 40, "O": 20, "P": 12})
    for row in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 120


def sheet_per_person(book, round_no: int) -> None:
    codes = coding(round_no)
    if not codes:
        return
    ids = sorted(next(iter(codes.values())))
    sheet = book.create_sheet(f"สรุปรายบุคคล รอบ {round_no}")
    header(sheet, ["ผู้ประเมิน", "จำนวนคำถาม", "คำตอบที่ถูกต้อง (รหัส 1, 2)",
                   "คำตอบที่ไม่ถูกต้อง (รหัส 0, 3)", "ค่าความถูกต้อง",
                   "ผ่านเกณฑ์ 80%"])
    total = 0
    for n in sorted(codes):
        right = sum(1 for i in ids if codes[n][i]["code"] in CORRECT)
        total += right
        sheet.append([f"คนที่ {n}", len(ids), right, len(ids) - right,
                      round(right / len(ids) * 100, 2),
                      "ผ่าน" if right / len(ids) >= 0.8 else "ไม่ผ่าน"])
    tests = len(ids) * len(codes)
    sheet.append([])
    sheet.append(["รวมทุกคน", tests, total, tests - total,
                  round(total / tests * 100, 2),
                  "ผ่าน" if total / tests >= 0.8 else "ไม่ผ่าน"])
    sheet.append([])
    sheet.append(["เกณฑ์ที่ใช้ตัดสิน คือผู้ประเมินทุกคนต้องได้ไม่น้อยกว่า 80% "
                  "ไม่ใช่ค่าเฉลี่ยของทั้งสามคน"])

    spread = Counter(codes[n][i]["code"] for n in codes for i in ids)
    sheet.append([])
    sheet.append(["การกระจายตัวของรหัสผลลัพธ์"])
    for cell in sheet[sheet.max_row]:
        cell.font = Font(bold=True)
    for code in (1, 2, 0, 3):
        sheet.append([CODE_LABEL[code], spread[code],
                      round(spread[code] / tests * 100, 2)])

    disputed = [i for i in ids
                if len({codes[n][i]["code"] in CORRECT for n in codes}) > 1]
    sheet.append([])
    sheet.append(["ตัดสินตรงกันทั้งสามคน", f"{len(ids) - len(disputed)}/{len(ids)}"])
    sheet.append(["เคสที่เห็นไม่ตรงกัน", ", ".join(disputed) or "—"])
    wrong_all = [i for i in ids
                 if all(codes[n][i]["code"] not in CORRECT for n in codes)]
    sheet.append(["เคสที่ทั้งสามคนเห็นว่าไม่ถูก", ", ".join(wrong_all) or "—"])
    widths(sheet, {"A": 46, "B": 14, "C": 26, "D": 26, "E": 16, "F": 16})


def sheet_rounds(book) -> None:
    sheet = book.create_sheet("สรุปรายรอบ")
    available = rounds_available()
    header(sheet, ["รอบ"] + [f"คนที่ {n}" for n in (1, 2, 3)] + ["โดยรวม", "หมายเหตุ"])
    for round_no in available:
        codes = coding(round_no)
        ids = sorted(next(iter(codes.values())))
        per, total = [], 0
        for n in (1, 2, 3):
            if n in codes:
                right = sum(1 for i in ids if codes[n][i]["code"] in CORRECT)
                per.append(round(right / len(ids) * 100, 2))
                total += right
            else:
                per.append("ไม่ได้วัด")
        counted = sum(1 for n in (1, 2, 3) if n in codes)
        overall = round(total / (len(ids) * counted) * 100, 2) if counted == 3 else "—"
        note = "" if counted == 3 else f"วัดได้ {counted} คน จึงไม่คิดค่าโดยรวม"
        sheet.append([round_no] + per + [overall, note])
    sheet.append([])
    sheet.append(["รอบ 1 ถึง 6 ผู้ประเมินเปิดคลังข้อมูลเองทั้งไฟล์ "
                  "ตั้งแต่รอบ 7 อ่านแฟ้มที่แนบตัวบทของข้อที่คำตอบอ้างมาให้ "
                  "ค่าระหว่างสองวิธีจึงเทียบกันตรง ๆ ไม่ได้"])
    widths(sheet, {"A": 8, "B": 12, "C": 12, "D": 12, "E": 12, "F": 60})


def sheet_dataset(book) -> None:
    with open(CORPUS, encoding="utf-8") as fh:
        recs = [json.loads(line) for line in fh if line.strip()]
    sheet = book.create_sheet("ชุดข้อมูล")
    header(sheet, ["ฉบับ", "จำนวนข้อ/มาตรา", "หน่วย", "วันประกาศ",
                   "สถานะ", "แหล่งที่มา"])
    groups: dict[str, list[dict]] = {}
    for r in recs:
        groups.setdefault(r["act_full"], []).append(r)
    for name, rows in sorted(groups.items()):
        repealed = sum(1 for r in rows if r.get("superseded_by"))
        sheet.append([
            name, len(rows), rows[0].get("unit", ""),
            rows[0].get("published", ""),
            "ยกเลิกแล้ว" if repealed == len(rows) else "ใช้บังคับอยู่",
            rows[0].get("source_url", ""),
        ])
        for cell in sheet[sheet.max_row]:
            cell.alignment = WRAP
    sheet.append([])
    sheet.append([f"รวมทั้งหมด {len(recs)} ระเบียน"])
    widths(sheet, {"A": 62, "B": 16, "C": 10, "D": 14, "E": 14, "F": 52})


def sheet_codebook(book) -> None:
    sheet = book.create_sheet("เกณฑ์การให้รหัส")
    header(sheet, ["หัวข้อ", "รายละเอียด"])
    for row in [
        ("รหัส 1", "คำถามที่ใช้คำตรงกับตัวบท และคำตอบถูกต้อง"),
        ("รหัส 2", "คำถามที่ใช้ภาษาพูด และคำตอบถูกต้อง"),
        ("รหัส 0", "คำถามที่ใช้คำตรงกับตัวบท และคำตอบไม่ถูกต้อง"),
        ("รหัส 3", "คำถามที่ใช้ภาษาพูด และคำตอบไม่ถูกต้อง"),
        ("นับว่าไม่ถูกต้องเมื่อ",
         "อ้างข้อที่ไม่ได้เขียนเรื่องนั้น / อ้างหมวดของผู้ประกอบวิชาชีพผิดประเภท "
         "(ข้อบังคับ 2550 เขียนซ้ำสี่รอบ ครู ข้อ 5-9 ผู้บริหารสถานศึกษา 10-14 "
         "ผู้บริหารการศึกษา 15-19 ศึกษานิเทศก์ 20-24) / ชี้อนุข้อผิด / "
         "สลับคำว่า ต้อง กับ พึง / ทำคำว่า อาจ ให้เป็นข้อบังคับ / "
         "เติมเงื่อนไข กำหนดเวลา วันที่ ผู้มีอำนาจ หรือตัวเลขที่ตัวบทไม่ได้เขียน / "
         "กลับความหมายของตัวบท / แจกแจงไม่ครบเมื่อคำถามถามว่ามีอะไรบ้าง / "
         "ปฏิเสธคำถามที่คลังตอบได้"),
        ("นับว่าถูกต้องเมื่อ",
         "สาระถูกและข้อที่อ้างเป็นข้อที่มีข้อความนั้นจริง สำนวนไม่สวย "
         "ข้อความถูกตัดตามความยาวที่จำกัด หรือมีรายละเอียดถูกต้องเพิ่มมา "
         "ไม่นับเป็นข้อผิด และการปฏิเสธคำถามที่อยู่นอกคลังนับว่าถูก"),
        ("เกณฑ์ผ่าน", "ผู้ประเมินทุกคนต้องได้ค่าความถูกต้องไม่น้อยกว่า 80%"),
        ("การเก็บคำตอบ",
         "เก็บคำตอบครั้งเดียวต่อรอบ แล้วให้ผู้ประเมินทั้งสามคนอ่านชุดเดียวกัน "
         "เพราะผู้ให้บริการแบบจำลองไม่รองรับ seed เรียกสามครั้งที่ "
         "temperature 0 ได้คำตอบต่างกันสองแบบ ถ้าให้แต่ละคนถามเอง "
         "ระยะห่างระหว่างคนจะวัดผู้ให้บริการ ไม่ใช่วัดผู้ประเมิน"),
        ("ตัวบทที่คาดหวัง",
         "ไม่ได้เขียนขึ้นเอง ข้อที่ผู้ประเมินทั้งสามคนเห็นว่าถูก ใช้ข้อที่คำตอบอ้าง "
         "เป็นความคาดหวัง เพราะมีผู้อ่านอิสระสามคนตรวจกับตัวบทแล้ว "
         "ข้อที่ไม่เป็นเช่นนั้น ใช้ข้อที่ปรากฏในบันทึกเหตุผลของผู้ประเมิน "
         "และเว้นว่างไว้เมื่อบันทึกไม่ได้ระบุข้อใด"),
    ]:
        sheet.append(list(row))
        for cell in sheet[sheet.max_row]:
            cell.alignment = WRAP
    widths(sheet, {"A": 24, "B": 110})


def main() -> None:
    round_no = int(sys.argv[1]) if len(sys.argv) > 1 else max(rounds_available())
    book = Workbook()
    book.remove(book.active)
    sheet_results(book, round_no)
    sheet_per_person(book, round_no)
    sheet_rounds(book)
    sheet_dataset(book)
    sheet_codebook(book)
    path = os.path.join(EVAL, f"ผลการทดสอบแชทบอท รอบ {round_no}.xlsx")
    book.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
