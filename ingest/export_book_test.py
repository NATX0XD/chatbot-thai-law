# -*- coding: utf-8 -*-
"""The networking bot's answers, laid out as the chatbot test sheet.

    .venv/bin/python -m ingest.check_book
    QUESTIONS=data/eval/network_exercise_questions.jsonl .venv/bin/python -m ingest.check_book
    .venv/bin/python -m ingest.export_book_test

Reads the two answer files check_book wrote and writes
data/eval/แบบทดสอบแชทบอทเครือข่าย.xlsx in the layout of the owner's Google
Sheet: chapter, question, answer, then two score columns.

The score columns are written EMPTY. A score is an assessor's reading of the
answer against the book, and nobody has done that reading for this dataset;
a number put there by this script would be a made-up result. The totals beside
the table are formulas over those columns and stay at zero until someone fills
them in.
"""
from __future__ import annotations

import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
EVAL = os.path.join(BASE_DIR, "data", "eval")
SOURCES = (os.path.join(EVAL, "network_exercise_answers.json"),
           os.path.join(EVAL, "network_answers.json"))
OUT_PATH = os.path.join(EVAL, "แบบทดสอบแชทบอทเครือข่าย.xlsx")

PINK = PatternFill("solid", fgColor="F4C7C3")
GREEN = PatternFill("solid", fgColor="D9EAD3")
TEAL = PatternFill("solid", fgColor="C9DDE0")
THIN = Side(style="thin", color="000000")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")
CHAPTER = re.compile(r"บทที่ (\d+)")

RUBRIC = ((2, "คำตอบถูกต้องและครบถ้วนตามเนื้อหาในตำรา"),
          (1, "คำตอบถูกต้องบางส่วน แต่ข้อมูลไม่ครบ"),
          (0, "คำตอบไม่ถูกต้องหรือไม่มีข้อมูลสนับสนุนจากตำรา"))


def load() -> list[dict]:
    rows = []
    for path in SOURCES:
        if not os.path.exists(path):
            raise SystemExit(f"ไม่พบ {path} -- รัน python -m ingest.check_book ก่อน")
        with open(path, encoding="utf-8") as handle:
            rows += json.load(handle)
    return rows


def chapter_of(row: dict):
    """The chapter the exercise is from, or the first one the answer cites."""
    if row.get("chapter"):
        return row["chapter"]
    cited = CHAPTER.search(" ".join(row.get("citations") or []) or row["answer"])
    return int(cited.group(1)) if cited and row["source"] == "book" else ""


def table(sheet, rows: list[dict]) -> int:
    """Write the question table and return the number of its last row."""
    heads = ("บทที่", "คำถามที่ใช้", "คำตอบที่ได้", "ประเมินความถูกต้อง", "ประเมิน RAG")
    for col, head in enumerate(heads, start=1):
        cell = sheet.cell(row=1, column=col, value=head)
        cell.font, cell.border = Font(bold=True), BOX
    sheet["D1"].fill, sheet["E1"].fill = PINK, GREEN
    for at, row in enumerate(rows, start=2):
        sheet.cell(row=at, column=1, value=chapter_of(row))
        sheet.cell(row=at, column=2, value=row["question"])
        sheet.cell(row=at, column=3, value=row["answer"])
        sheet.cell(row=at, column=4).fill = PINK
        sheet.cell(row=at, column=5).fill = GREEN
        for col in range(1, 6):
            sheet.cell(row=at, column=col).alignment = WRAP
            sheet.cell(row=at, column=col).border = BOX
    last = len(rows) + 1
    for col, allowed in (("D", '"0,1"'), ("E", '"0,1,2"')):
        rule = DataValidation(type="list", formula1=allowed, allow_blank=True)
        sheet.add_data_validation(rule)
        rule.add(f"{col}2:{col}{last}")
    for col, width in zip("ABCDE", (8, 46, 90, 20, 14)):
        sheet.column_dimensions[col].width = width
    sheet.freeze_panes = "A2"
    return last


def side_tables(sheet, last: int) -> None:
    """The accuracy count and the scoring rubric, to the right as in the sample."""
    sheet["H1"], sheet["I1"] = "คำถามตรง ตอบตรง", "คำถามตรง ตอบไม่ตรง"
    sheet["G2"] = "ประเมินค่าความถูกต้อง"
    sheet["H2"] = f"=COUNTIF(D2:D{last},1)"
    sheet["I2"] = f"=COUNTIF(D2:D{last},0)"
    for ref in ("H1", "I1", "G2", "H2", "I2"):
        sheet[ref].fill, sheet[ref].border = PINK, BOX
    sheet["K2"] = "ประเมิน RAG and LLM"
    sheet["K2"].fill = TEAL
    for col, head in zip("LMN", ("ระดับ", "เกณฑ์", "คะแนน")):
        sheet[f"{col}1"] = head
        sheet[f"{col}1"].font = Font(bold=True)
        sheet[f"{col}1"].fill, sheet[f"{col}1"].border = TEAL, BOX
    for at, (level, wording) in enumerate(RUBRIC, start=2):
        sheet[f"L{at}"], sheet[f"M{at}"] = level, wording
        sheet[f"N{at}"] = level
        for col in "LMN":
            sheet[f"{col}{at}"].fill, sheet[f"{col}{at}"].border = TEAL, BOX
    for col, width in zip("GHIKLMN", (24, 18, 20, 22, 8, 52, 10)):
        sheet.column_dimensions[col].width = width


def main() -> None:
    rows = load()
    inside = [r for r in rows if r["group"] == "book"]
    outside = [r for r in rows if r["group"] != "book"]
    inside.sort(key=lambda r: (chapter_of(r) or 99))
    book = Workbook()
    sheet = book.active
    sheet.title = "คำถามจากหนังสือ"
    side_tables(sheet, table(sheet, inside))
    other = book.create_sheet("คำถามนอกหนังสือ")
    side_tables(other, table(other, outside))
    book.save(OUT_PATH)
    print(f"{len(inside)} ข้อจากหนังสือ, {len(outside)} ข้อนอกหนังสือ -> {OUT_PATH}")
    print("ช่องคะแนนเว้นว่างไว้ให้ผู้ประเมินกรอก")


if __name__ == "__main__":
    main()
